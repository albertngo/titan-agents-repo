"""HikePOS update (Albert, 2026-10-07): existing rows keep Name and SKU, new products take the
Airtable SKU, and Hike keeps its own +$1.40 markup on the Airtable cost."""

import json
import sys
import tempfile
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))

import hike_update as hu  # noqa: E402

REG = json.loads((Path(__file__).resolve().parent.parent / "platform-settings" / "hike.json").read_text())
CFG = REG["suppliers"]["FAW"]
OUT = REG["outlet"]


def at(sku, name, coll="", box=None, width=None, cost=1.99, active=True, **kw):
    return {"sku": sku, "name": name, "coll": coll, "box": box, "width": width, "cost": cost,
            "active": active, "eff": "2026-09-19", **kw}


def hike(sku, name, cost=2.29, retail=3.69, active="TRUE", desc="3/31/25 -  -  / Promotion: $ / 2.99=70c"):
    return {"SKU": sku, "Barcode": sku, "Name": name, "Description": desc, "Active": active,
            f"{OUT}_Cost price": cost, f"{OUT}_Retail price": retail, f"{OUT}_Price Excluding Tax": retail}


class TestParse(unittest.TestCase):
    def test_reads_category_colour_box_and_width(self):
        p = hu.parse_hike_name('(C) FAWENG - NAF - Hickory (Bronze) T&G | #FAW.E.H.Bro.6.1 |  - 6.5" x 18mm x 1800mm (80%) - 19.18sf/b')
        self.assertEqual((p["cat"], p["colour"], p["box"], p["width"]), ("ENG", "Bronze", 19.18, 6.5))

    def test_fractional_width(self):
        p = hu.parse_hike_name('(C) FAWHAR - NAF - Walnut (Tan) T&G | #FAW.H.EW.Tan.4.1 |  - 4 3/4" x 3/4" x RL - 22.06sf/b')
        self.assertEqual(p["width"], 4.75)


class TestMatch(unittest.TestCase):
    def test_box_size_picks_between_two_collections_with_one_colour(self):
        recs = [at("LVP-FAWK-0001", "NAF 6.5mm SPC Vinyl 7.1\" — Windsor", box=23.9, width=7.1),
                at("LVP-FAWK-0046", "NAF 8mm Royal Flooring Vinyl 7.1\" — Windsor", box=14.93, width=7.1)]
        rows = [hike("FAW.V.SW.Win.8", '(C) FAWVIN - Royal Flooring Solutions - SPC w/underpad (Windsor) Click | #FAW.V.SW.Win.8 |  - 8mm x 182mm x 1524mm - 14.93sf/b')]
        m, _, _ = hu.match(rows, recs, CFG)
        self.assertEqual(m, {0: "LVP-FAWK-0046"})

    def test_a_5_inch_product_never_takes_a_6_5_inch_price(self):
        recs = [at("ENG-FAWK-0006", "NAF Engineered Hickory 6.5\" — Downtown Grey", box=19.18, width=6.5)]
        rows = [hike("FAW.E.H.DowGre.5.1", '(C) FAWENG - Hickory (Downtown Grey) T&G | #FAW.E.H.DowGre.5.1 |  - 5" x 12mm x RL - 26.25sf/b')]
        m, _, notes = hu.match(rows, recs, CFG)
        self.assertEqual(m, {})
        self.assertIn("width", notes[0][0])

    def test_species_code_separates_hickory_maple_and_walnut_natural(self):
        recs = [at("ENG-FAWK-0075", "NAF Engineered Exotic Walnut 5\" — Natural", box=19.68, width=5),
                at("ENG-FAWK-0077", "NAF Engineered Maple 6.5\" — Natural", box=19.18, width=6.5)]
        rows = [hike("FAW.E.H.Nat.6.1", '(C) FAWENG - NAF - Hickory (Natural) T&G | #FAW.E.H.Nat.6.1 |  - 6.5" x 18mm - 20.25sf/b'),
                hike("FAW.E.M.Nat.5.1", '(C) FAWENG - NAF - Maple (Natural) T&G | #FAW.E.M.Nat.5.1 |  - 6.5" x 18mm x RL - 19.18sf/b')]
        m, _, _ = hu.match(rows, recs, CFG)
        self.assertEqual(m, {1: "ENG-FAWK-0077"})

    def test_spelling_alias_and_hand_match_come_from_the_registry(self):
        recs = [at("ENG-FAWK-0042", "NAF Engineered White Oak Regal 7.5\" — Bermuda", box=23.31, width=7.5),
                at("ENG-FAWK-0022", "NAF Engineered White Oak Click 5\" — Golden", box=29.28, width=5)]
        rows = [hike("FAW.E.O.Bem.7", '(C) FAWENG - NAF - Regal - Oak (Bemuda) T&G | #FAW.E.O.Bem.7 |  - 7.5" x 18mm - 23.31sf/b'),
                hike("100332", "(C) FAWENG - GOLDEN - Click 29.28sf/b")]
        m, _, _ = hu.match(rows, recs, CFG)
        self.assertEqual(m, {0: "ENG-FAWK-0042", 1: "ENG-FAWK-0022"})

    def test_primary_row_is_the_active_one_with_the_right_box(self):
        recs = [at("LVP-FAWK-0008", "NAF 7mm Aquaplus Select Vinyl 7.2\" — Deerhurst", box=17.91, width=7.2)]
        rows = [hike("FAW.V.SW.Dee.7", '(C) FAWVIN - Select (Deerhurst) Click | #FAW.V.SW.Dee.7 |  - 7mm X 180mm (7.2" wide) - 14.76sf/b'),
                hike("FAW.V.7V.Dee.7", '(C) FAWVIN - Select (Deerhurst) Click | #FAW.V.7V.Dee.7 |  - 7mm X 182mm (7.2" wide) - 17.91sf/b')]
        m, primary, _ = hu.match(rows, recs, CFG)
        self.assertEqual(len(m), 2)
        self.assertEqual(primary["LVP-FAWK-0008"], 1)


class TestValues(unittest.TestCase):
    def test_flooring_retail_is_cost_plus_the_hike_markup(self):
        self.assertEqual(hu.new_retail(at("LVP-FAWK-0011", "x"), 1.99, 2.29, 3.69, 1.40), 3.39)

    def test_an_accessory_keeps_its_own_markup(self):
        self.assertEqual(hu.new_retail(at("ACC-FAWK-0001", "x"), 49, 49, 71.5, 1.40), 71.5)

    def test_an_ended_promo_is_not_the_cost(self):
        o = at("X", "x", cost=3.69, promo=3.29, promoend="2026-09-30")
        self.assertEqual(hu.effective_cost(o, "2026-10-07"), 3.69)
        self.assertEqual(hu.effective_cost(o, "2026-09-20"), 3.29)

    def test_description_date_and_70c_floor_follow_the_new_cost(self):
        d = "3/31/25 -  - Husssh Gold only / Promotion: $ / 4.39=70c"
        self.assertEqual(hu.rewrite_description(d, "2026-09-19", 3.69, 3.79, 0.70),
                         "2026-09-19 -  - Husssh Gold only / Promotion: $ / 4.49=70c")

    def test_a_floor_that_is_not_cost_plus_70c_is_left_alone(self):
        d = "3/31/25 -  -  / Promotion: $ / 30=70c"
        self.assertEqual(hu.rewrite_description(d, None, 18, 18, 0.70), d)

    def test_new_name_comes_from_lightspeed_without_the_promo_marker(self):
        name = hu.hike_name_for_new(at("LVP-FAWK-0013", "x"), CFG, {
            "LVP-FAWK-0013": "(P 2026-09-30) NAFLVP-SPC - Aquaplus Select (Grand Bend) Click | 7.2\" x 7mm x RL - 17.91sf/b"})
        self.assertEqual(name, '(C) FAWVIN - NAF - Aquaplus Select (Grand Bend) Click | #LVP-FAWK-0013 |  - 7.2" x 7mm x RL - 17.91sf/b')


class TestWorkbook(unittest.TestCase):
    def test_existing_name_and_sku_never_change_and_new_rows_use_the_airtable_sku(self):
        from openpyxl import Workbook, load_workbook
        header = ["Name", "Description", "SKU", "Barcode", "Product type", "Product tag", "Supplier name", "Slug",
                  "Image URL", "Active", "Track inventory", "Allow out of stock", f"{OUT}_Cost price",
                  f"{OUT}_Retail price", f"{OUT}_Average Cost Price", f"{OUT}_Stock", f"{OUT}_Stock on hand",
                  f"{OUT}_Reorder level", f"{OUT}_Reorder value", f"{OUT}_Price Excluding Tax"]
        name = '(C) FAWVIN - NAF AQUAPLUS - Select - SPC w/underpad (Wasaga) Click | #FAW.V.SW.Was.7 |  - 7mm X 182mm X 1524mm, 7.2" wide - 17.91sf/b'
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "in.xlsx"
            wb = Workbook(); ws = wb.active; ws.title = "Product"; ws.append(header)
            ws.append([name, "3/31/25 -  -  / Promotion: $ / 2.99=70c", "FAW.V.SW.Was.7", "FAW.V.SW.Was.7", "VINYL",
                       "Flooring;C", "FAW", "slug", None, "TRUE", "FALSE", "TRUE", 2.29, 3.69, 2.29, 12, 12, None, None, 3.69])
            wb.save(src)
            wb, ws, hdr, rows = hu.read_hike(src)
            recs = [at("LVP-FAWK-0011", "NAF 7mm Aquaplus Select Vinyl 7.2\" — Wasaga", coll="Aquaplus Select", box=17.91, width=7.2),
                    at("LVP-FAWK-0013", "NAF 7mm Aquaplus Select Vinyl 7.2\" — Grand Bend", coll="Aquaplus Select", box=17.91, width=7.2)]
            upd, new, rep = hu.build(rows, hdr, recs, CFG, OUT, {}, "2026-10-07")
            out = Path(d) / "out.xlsx"
            hu.write_outputs(wb, ws, hdr, rows, upd, new, out, rep, Path(d) / "r.csv",
                             set(REG["never_change_on_existing"]) | set(REG["never_touch_columns"]))
            got = load_workbook(out).active
            r2 = dict(zip(hdr, [c.value for c in got[2]]))
            self.assertEqual((r2["Name"], r2["SKU"], r2["Barcode"], r2[f"{OUT}_Stock"]), (name, "FAW.V.SW.Was.7", "FAW.V.SW.Was.7", 12))
            self.assertEqual((r2[f"{OUT}_Cost price"], r2[f"{OUT}_Retail price"]), (1.99, 3.39))
            r3 = dict(zip(hdr, [c.value for c in got[3]]))
            self.assertEqual((r3["SKU"], r3["Barcode"], r3[f"{OUT}_Retail price"]), ("LVP-FAWK-0013", "LVP-FAWK-0013", 3.39))
            self.assertEqual(r3["Description"], "2026-09-19 -  -  / Promotion: $ / 2.69=70c")


if __name__ == "__main__":
    unittest.main()
