#!/usr/bin/env python3
"""Tests for the read-only Lightspeed client and pull script.

Stdlib unittest, no pytest — matching the repo's stdlib-first script convention.
`openpyxl` is needed only by the compare-export cases, which skip without it.

    python3 -m unittest discover -s tests -v

The pagination cases matter more than they look. The X-Series response envelope
and cursor semantics are CONFIGURED in platform-settings/lightspeed.json, not yet
verified against the live account (run `lightspeed_pull.py --probe` for that), so
the walk is written to survive being wrong in either direction — an inclusive
cursor, or a server that stops advancing. Getting that wrong duplicates or drops
products silently rather than raising.
"""

import importlib.util
import json
import os
import sys
import unittest
from datetime import datetime, timezone
from pathlib import Path
from unittest import mock

REPO_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(REPO_ROOT / "scripts"))


def _load(name, path):
    """Load a script as a module, registering it in sys.modules first.

    Registration matters: lightspeed_write.py does `from lightspeed_client import
    LightspeedError`, and without this the test would hold a different module
    object than the writer does — so `assertRaises(LightspeedError)` would not
    catch the writer's own exceptions. Same class, two identities.
    """
    if name in sys.modules:
        return sys.modules[name]
    spec = importlib.util.spec_from_file_location(name, REPO_ROOT / path)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


lsc = _load("lightspeed_client", "scripts/lightspeed_client.py")
lp = _load("lightspeed_pull", "scripts/lightspeed_pull.py")
lw = _load("lightspeed_write", "scripts/lightspeed_write.py")

GRANDEUR = REPO_ROOT / "ingest/2026-09-03/grandeur_ls_product_export_2026-09-03.xlsx"
CANADIAN = REPO_ROOT / "ingest/2026-09-03/canadian_standard_ls_product_export_2026-09-03.xlsx"


def stub_client(pages, page_size):
    """A client whose GET returns canned pages instead of touching the network."""
    cl = lsc.LightspeedClient(domain_prefix="x", token="y", config=lsc.load_config())
    cl.min_interval = 0
    cl.page_size = page_size
    it = iter(pages)

    def get(path, params=None):
        try:
            return next(it)
        except StopIteration:
            raise AssertionError("the walk requested more pages than the stub provides")

    cl.get = get
    return cl


def ids(records):
    return [r["id"] for r in records]


class TestNoWritePath(unittest.TestCase):
    """The Phase 1 guarantee is structural: neither file can write to Lightspeed."""

    def test_no_write_verbs(self):
        for rel in ("scripts/lightspeed_client.py", "scripts/lightspeed_pull.py"):
            src = (REPO_ROOT / rel).read_text()
            for verb in ("POST", "PUT", "PATCH", "DELETE"):
                self.assertNotIn(f'"{verb}"', src,
                                 f"{rel} contains a {verb} literal — the read-only "
                                 "guarantee in its docstring is no longer true")


class TestWriter(unittest.TestCase):
    """The only module that can change the POS. Its guardrails are the tests."""

    def writer(self, dry_run=False, capture=None):
        w = lw.LightspeedWriter(domain_prefix="x", token="y",
                                config=lsc.load_config(), dry_run=dry_run)
        w.min_interval = 0
        if capture is not None:
            # stand in for the network at the lowest level, so anything that
            # reaches the wire is visible to the test
            def fake(url, method, body=None):
                capture.append({"method": method, "url": url, "body": body})
                raise AssertionError("a live request was attempted")
            w._open = fake
        return w

    def test_dry_run_never_builds_a_request(self):
        sent = []
        w = self.writer(dry_run=True, capture=sent)
        w.create_family({"name": "T", "variants": [{"sku": "A-1"}]})
        w.update_variant("id-1", {"supply_price": 1.0})
        self.assertEqual(sent, [], "dry run reached the network")
        self.assertEqual(len(w.planned), 2)
        self.assertEqual(w.write_stats()["writes_sent"], 0)

    def test_dry_run_returns_no_usable_id(self):
        """A placeholder id would get written into Airtable as if it were real."""
        w = self.writer(dry_run=True)
        self.assertIsNone(w.create_family({"name": "T", "variants": [{"sku": "A-1"}]}))

    def test_create_requires_a_name(self):
        w = self.writer(dry_run=True)
        with self.assertRaises(lsc.LightspeedError):
            w.create_family({"variants": [{"sku": "A-1"}]})

    def test_every_variant_needs_an_explicit_sku(self):
        """Lightspeed mints one otherwise, which RULE 0 forbids."""
        w = self.writer(dry_run=True)
        with self.assertRaises(lsc.LightspeedError) as cm:
            w.create_family({"name": "T", "variants": [{"sku": "A-1"}, {"supply_price": 2}]})
        self.assertIn("sku", str(cm.exception))

    def test_create_returns_raw_ids_not_a_mapping(self):
        """Pairing ids to SKUs positionally is the Grandeur mistake."""
        w = self.writer()
        w._send = lambda m, p, params=None, body=None: {"data": ["id-a", "id-b"]}
        got = w.create_family({"name": "T", "variants": [{"sku": "A-1"}, {"sku": "A-2"}]})
        self.assertEqual(got, ["id-a", "id-b"])
        self.assertNotIsInstance(got, dict)

    def test_unexpected_create_response_raises(self):
        w = self.writer()
        w._send = lambda m, p, params=None, body=None: {"data": {}}
        with self.assertRaises(lsc.LightspeedError):
            w.create_family({"name": "T", "variants": [{"sku": "A-1"}]})

    def test_common_section_is_refused_without_a_reason(self):
        """`common` rewrites every family member, and name regroups families."""
        w = self.writer(dry_run=True)
        with self.assertRaises(lsc.LightspeedError) as cm:
            w.update_variant("id-1", {"supply_price": 1.0}, common={"name": "New"})
        self.assertIn("common", str(cm.exception))
        w.update_variant("id-1", {"supply_price": 1.0}, common={"name": "New"},
                         allow_common_reason="explicit operator intent")

    def test_update_with_nothing_to_write_raises(self):
        w = self.writer(dry_run=True)
        with self.assertRaises(lsc.LightspeedError):
            w.update_variant("id-1", {})

    def test_a_pure_rename_sends_common_only(self):
        """Renaming a family has no per-variant details, and that is legitimate."""
        w = self.writer()
        seen = {}
        w._send = lambda m, p, params=None, body=None: (seen.update(body=body), {})[1]
        w.update_variant("id-1", {}, common={"name": "New Family Name"},
                         allow_common_reason="converting a singleton into a variant group")
        self.assertEqual(seen["body"], {"common": {"name": "New Family Name"}})
        self.assertNotIn("details", seen["body"])

    def test_update_payload_shape(self):
        w = self.writer()
        seen = {}
        def cap(m, p, params=None, body=None):
            seen.update(method=m, path=p, body=body); return {}
        w._send = cap
        w.update_variant("id-1", {"supply_price": 3.5, "price_excluding_tax": 4.5})
        self.assertEqual(seen["method"], "PUT")
        self.assertIn("2.1", seen["path"])
        self.assertEqual(seen["body"], {"details": {"supply_price": 3.5,
                                                    "price_excluding_tax": 4.5}})
        self.assertNotIn("common", seen["body"])

    def test_add_variant_targets_the_2_1_collection_and_joins_by_name(self):
        w = self.writer()
        seen = {}
        w._send = lambda m, p, params=None, body=None: (seen.update(method=m, path=p, body=body), {})[1]
        w.add_variant("VIDENG - 7 AWO (Snowwhite) T&G", "ENG-VIDR-0195",
                      [{"attribute_id": "grade-id", "attribute_value": "Character"}],
                      details={"price_excluding_tax": 5.79})
        self.assertEqual(seen["method"], "POST")
        self.assertEqual(seen["path"], "/api/2.1/products/")
        self.assertEqual(seen["body"]["common"], {"name": "VIDENG - 7 AWO (Snowwhite) T&G"})
        d = seen["body"]["details"]
        self.assertEqual(d["product_codes"], [{"code": "ENG-VIDR-0195", "type": "CUSTOM"}])
        self.assertEqual(d["variant_attribute_values"][0]["attribute_value"], "Character")
        self.assertEqual(d["price_excluding_tax"], 5.79)

    def test_add_variant_refuses_incomplete_input(self):
        w = self.writer(dry_run=True)
        vals = [{"attribute_id": "g", "attribute_value": "Character"}]
        for args in ((None, "SKU-1", vals), ("Fam", None, vals), ("Fam", "SKU-1", [])):
            with self.subTest(args=args), self.assertRaises(lsc.LightspeedError):
                w.add_variant(*args)

    def test_add_variant_is_intercepted_by_dry_run(self):
        sent = []
        w = self.writer(dry_run=True, capture=sent)
        w.add_variant("Fam", "SKU-1", [{"attribute_id": "g", "attribute_value": "C"}])
        self.assertEqual(sent, [])
        self.assertEqual(len(w.planned), 1)

    def test_family_read_uses_the_3_0_key_names(self):
        """Three names for one concept; reading the wrong one fails silently."""
        w = self.writer()
        fam = {"variants": [{"id": "v1", "primary_sku_code": "ENG-VIDR-0038",
                             "product_codes": [{"type": "CUSTOM", "code": "ENG-VIDR-0038"}],
                             "variant_definitions": [{"attribute_id": "g", "name": "Grade",
                                                      "value": "Select"}]}]}
        w.read_family = lambda pid: fam
        self.assertEqual(list(w.family_by_sku("x")), ["ENG-VIDR-0038"])
        vals = w.family_attribute_values("x")
        self.assertEqual(vals["v1"], [{"attribute_id": "g", "name": "Grade", "value": "Select"}])

    def test_family_attribute_values_survives_the_2_0_spelling(self):
        w = self.writer()
        w.read_family = lambda pid: {"variants": [
            {"id": "v1", "variant_options": [{"attribute_id": "g", "name": "Grade",
                                              "value": "Character"}]}]}
        self.assertEqual(w.family_attribute_values("x")["v1"][0]["value"], "Character")

    def test_there_is_no_delete_capability(self):
        w = self.writer(dry_run=True)
        for banned in ("delete", "delete_product", "deactivate", "archive"):
            self.assertFalse(hasattr(w, banned), f"writer exposes {banned}")
        src = (REPO_ROOT / "scripts/lightspeed_write.py").read_text()
        self.assertNotIn('"DELETE"', src.replace('WRITE_METHODS = ("POST", "PUT", "PATCH", "DELETE")', ''))

    def test_reads_still_go_through_the_read_path(self):
        """read_family must not be caught by the dry-run write interceptor."""
        w = self.writer(dry_run=True)
        calls = []

        class FakeResponse:
            def read(self): return b'{"data": {"variants": []}}'
            def __enter__(self): return self
            def __exit__(self, *a): return False

        w._open = lambda url, method, body=None: (calls.append(method), FakeResponse())[1]
        w.read_family("id-1")
        self.assertEqual(calls, ["GET"])
        self.assertEqual(w.planned, [], "a read was recorded as a planned write")


class TestRetryAfter(unittest.TestCase):
    NOW = datetime(2026, 9, 10, 12, 0, 0, tzinfo=timezone.utc)

    def test_forms(self):
        cases = [
            ("30", 30.0),
            ("0", 0.0),
            ("", None),
            (None, None),
            ("garbage", None),
            ("Thu, 10 Sep 2026 12:00:45 GMT", 45.0),   # the RFC 1123 form LS documents
            ("Thu, 10 Sep 2026 11:59:00 GMT", 0.0),    # already elapsed -> clamp, never negative
        ]
        for raw, want in cases:
            with self.subTest(raw=raw):
                self.assertEqual(lsc._retry_after_seconds(raw, now=self.NOW), want)


class TestPagination(unittest.TestCase):

    def test_version_block_cursor(self):
        pages = [
            {"data": [{"id": f"a{i}", "version": 100 + i} for i in range(3)],
             "version": {"max": 102}},
            {"data": [{"id": "b0", "version": 103}], "version": {"max": 103}},
        ]
        self.assertEqual(ids(stub_client(pages, 3).paginate("/p")), ["a0", "a1", "a2", "b0"])

    def test_bare_list_envelope(self):
        pages = [[{"id": "a", "version": 5}, {"id": "b", "version": 9}],
                 [{"id": "c", "version": 11}]]
        self.assertEqual(ids(stub_client(pages, 2).paginate("/p")), ["a", "b", "c"])

    def test_stuck_cursor_stops_without_duplicating(self):
        """A server that keeps returning the same page must not re-emit it."""
        page = {"data": [{"id": "a", "version": 7}], "version": {"max": 7}}
        self.assertEqual(ids(stub_client([page, page], 1).paginate("/p")), ["a"])

    def test_inclusive_cursor_does_not_duplicate_the_boundary(self):
        """`after` is documented exclusive; stay correct if it is ever inclusive."""
        pages = [
            {"data": [{"id": "a", "version": 1}, {"id": "b", "version": 2}], "version": {"max": 2}},
            {"data": [{"id": "b", "version": 2}, {"id": "c", "version": 3}], "version": {"max": 3}},
            {"data": [], "version": {"max": 3}},
        ]
        self.assertEqual(ids(stub_client(pages, 2).paginate("/p")), ["a", "b", "c"])

    def test_short_page_terminates(self):
        pages = [{"data": [{"id": "a", "version": 1}], "version": {"max": 1}}]
        self.assertEqual(ids(stub_client(pages, 5).paginate("/p")), ["a"])

    def test_empty_page(self):
        self.assertEqual(list(stub_client([{"data": []}], 5).paginate("/p")), [])

    def test_max_pages(self):
        pages = [{"data": [{"id": "a", "version": 1}], "version": {"max": 1}},
                 {"data": [{"id": "b", "version": 2}], "version": {"max": 2}}]
        self.assertEqual(ids(stub_client(pages, 1).paginate("/p", max_pages=1)), ["a"])

    def test_unknown_envelope_raises(self):
        with self.assertRaises(lsc.LightspeedError):
            list(stub_client([{"data": {"nope": 1}}], 5).paginate("/p"))

    def test_no_derivable_cursor_raises(self):
        with self.assertRaises(lsc.LightspeedError):
            list(stub_client([{"data": [{"id": "a"}]}], 5).paginate("/p"))


class TestCredentials(unittest.TestCase):

    def test_missing_credentials_names_the_variable(self):
        # The client falls back to os.environ, so this must run with the real
        # variables cleared — otherwise it silently stops testing anything the
        # moment a developer has credentials configured, which is everyone now.
        with mock.patch.dict(os.environ, {}, clear=True):
            with self.assertRaises(lsc.LightspeedError) as cm:
                lsc.LightspeedClient(domain_prefix=None, token=None, config=lsc.load_config())
        self.assertIn("LIGHTSPEED_DOMAIN_PREFIX", str(cm.exception))

    def test_env_is_used_when_arguments_are_omitted(self):
        with mock.patch.dict(os.environ,
                             {"LIGHTSPEED_DOMAIN_PREFIX": "envprefix",
                              "LIGHTSPEED_PERSONAL_TOKEN": "envtoken"}, clear=True):
            c = lsc.LightspeedClient(config=lsc.load_config())
        self.assertEqual(c.base_url, "https://envprefix.retail.lightspeed.app")

    def test_api_version_override(self):
        with mock.patch.dict(os.environ,
                             {"LIGHTSPEED_DOMAIN_PREFIX": "p", "LIGHTSPEED_PERSONAL_TOKEN": "t",
                              "LIGHTSPEED_API_VERSION": "2026-01"}, clear=True):
            self.assertEqual(lsc.LightspeedClient(config=lsc.load_config()).api_version, "2026-01")


class TestSupplierExtraction(unittest.TestCase):

    def test_supplier_paths(self):
        self.assertEqual(lp.supplier_of({"supplier_name": " Lee Flooring "}), "Lee Flooring")
        self.assertEqual(lp.supplier_of({"supplier": {"name": "Grandeur"}}), "Grandeur")
        self.assertEqual(lp.supplier_of({"supplier": "Vidar"}), "Vidar")
        self.assertIsNone(lp.supplier_of({"supplier": {"id": "abc"}}))
        self.assertIsNone(lp.supplier_of({"supplier_name": "   "}))
        self.assertIsNone(lp.supplier_of({}))

    def test_supplier_falls_back_to_product_suppliers(self):
        """6 of 14,525 live products carry a supplier only here."""
        rec = {"supplier": None,
               "product_suppliers": [{"supplier_id": "x", "supplier_name": "WINCA / OLYMPIA"}]}
        self.assertEqual(lp.supplier_of(rec), "WINCA / OLYMPIA")
        self.assertIsNone(lp.supplier_of({"supplier": None, "product_suppliers": []}))
        self.assertIsNone(lp.supplier_of({"supplier": None,
                                          "product_suppliers": [{"supplier_id": None}]}))

    def test_category_is_an_object_on_the_api(self):
        """product_category comes back as {name: ...}, leaf name only, or null."""
        self.assertEqual(lp.category_of({"product_category": {"name": "SPC"}}), "SPC")
        self.assertEqual(lp.category_of({"product_category": "TILE"}), "TILE")
        self.assertIsNone(lp.category_of({"product_category": None}))
        self.assertIsNone(lp.category_of({"product_category": {"id": "x"}}))
        self.assertIsNone(lp.category_of({}))

    def test_dig_survives_non_dicts(self):
        self.assertIsNone(lp.dig({"a": "scalar"}, ("a", "b")))
        self.assertIsNone(lp.dig(None, ("a",)))


class TestRegistry(unittest.TestCase):
    """The registry is load-bearing config — a typo there breaks every run."""

    def test_lightspeed_registry_has_what_the_client_reads(self):
        cfg = lsc.load_config()
        api, pg, rl = cfg["api"], cfg["api"]["pagination"], cfg["rate_limits"]
        self.assertIn("{domain_prefix}", api["base_url_template"])
        self.assertTrue(api["endpoints"]["products"].startswith("/"))
        for k in ("cursor_param", "page_size_param", "page_size", "cursor_field"):
            self.assertIn(k, pg)
        for k in ("retry_after_header", "limiter_type_header", "min_interval_ms", "max_retries"):
            self.assertIn(k, rl)

    def test_leaf_categories_exclude_the_non_leaves(self):
        cats = lsc.load_config()["product_categories"]
        for bad in cats["not_leaves_never_write"]:
            self.assertNotIn(bad, cats["leaves"],
                             f"{bad} is listed as a leaf but LS rejects it as an "
                             "intermediate node")

    def test_supplier_aliases_are_unique_both_ways(self):
        aliases = json.loads(
            (REPO_ROOT / "platform-settings/airtable-destinations.json").read_text()
        )["supplier_aliases"]
        pairs = {k: v for k, v in aliases.items() if not k.startswith("_")}
        self.assertEqual(len(set(pairs.values())), len(pairs),
                         "two Notion companies map to the same Airtable supplier")


class TestCompareExport(unittest.TestCase):
    """Ground truth: the two LS exports a person hand-made for the 2026-09-03 backfill."""

    @classmethod
    def setUpClass(cls):
        try:
            from openpyxl import load_workbook  # noqa: F401
        except ImportError:
            raise unittest.SkipTest("openpyxl not installed")

    @staticmethod
    def rows_from(path):
        from openpyxl import load_workbook
        # Deliberately NOT read_only=True — see the note in compare_export().
        ws = load_workbook(path, data_only=True).active
        it = ws.iter_rows(values_only=True)
        header = [str(c).strip().lower() for c in next(it)]
        si, ii = header.index("sku"), header.index("id")
        return [{"sku": str(r[si]).strip(), "id": str(r[ii]).strip()}
                for r in it if r[si] not in (None, "")]

    def test_read_only_mode_would_misread_these_exports(self):
        """Guards the bug that read_only=True hides: 1 column instead of 33."""
        from openpyxl import load_workbook
        ws = load_workbook(GRANDEUR, read_only=True, data_only=True).active
        header = next(ws.iter_rows(values_only=True))
        self.assertEqual(len(header), 1,
                         "read_only no longer misreads this file — the warning in "
                         "compare_export() and the ls-id-backfill skill can be revisited")
        ws2 = load_workbook(GRANDEUR, data_only=True).active
        self.assertEqual(len(next(ws2.iter_rows(values_only=True))), 33)

    def test_identical_data_reconciles(self):
        for path, expected in ((GRANDEUR, 267), (CANADIAN, 311)):
            with self.subTest(path=path.name):
                pull = self.rows_from(path)
                self.assertEqual(len(pull), expected)
                self.assertEqual(lp.compare_export(pull, path), 0)

    def test_uuid_disagreement_is_caught(self):
        """The Grandeur failure mode: a UUID that belongs to a different product."""
        pull = self.rows_from(GRANDEUR)
        pull[0]["id"] = "00000000-dead-beef-0000-000000000000"
        self.assertEqual(lp.compare_export(pull, GRANDEUR), 1)

    def test_missing_sku_is_reported_but_not_fatal(self):
        """A product gone since the export is drift, not a reading error."""
        pull = self.rows_from(GRANDEUR)
        pull.pop(1)
        self.assertEqual(lp.compare_export(pull, GRANDEUR), 0)


if __name__ == "__main__":
    unittest.main()
