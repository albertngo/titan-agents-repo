#!/usr/bin/env python3
"""Pull the Lightspeed (X-Series) product catalogue. READ ONLY.

GET only — this script has no code path that writes to Lightspeed. It imports
scripts/lightspeed_client.py, which exposes GET and pagination and nothing else.
The write side lands separately with the catalogue-sync actions agents, so that
granting a write credential and shipping write code stay two decisions.

This replaces the manual step in the current backfill loop, where a person exports
the LS product list to .xlsx by hand before ls-id-backfill can run.

The products endpoint has no server-side supplier filter, so the walk fetches the
whole catalogue once, caches the raw records, and filters locally. The cache lives
under analysis/cache/ (gitignored, regenerable) so a re-run during a sync does not
re-page the catalogue and spend rate-limit budget for nothing. --refresh forces a
new walk.

Usage:
    # First run against a live account — dump the real response shape and stop.
    # Do this before trusting the pagination config; nothing is cached or written.
    python3 scripts/lightspeed_pull.py --probe

    # Full catalogue -> ingest/<today>/lightspeed-products.json
    python3 scripts/lightspeed_pull.py

    # One supplier, from cache if today's walk already happened
    python3 scripts/lightspeed_pull.py --supplier "Lee Flooring"

    # Compare a pull against a hand-exported LS workbook (the Phase 1 check)
    python3 scripts/lightspeed_pull.py --compare-export ingest/2026-09-03/grandeur_ls_product_export_2026-09-03.xlsx

Environment: LIGHTSPEED_DOMAIN_PREFIX, LIGHTSPEED_PERSONAL_TOKEN — see .env.example.
Config: platform-settings/lightspeed.json.
"""

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

sys.path.insert(0, str(Path(__file__).resolve().parent))
from lightspeed_client import LightspeedClient, LightspeedError, load_config  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
CACHE_DIR = REPO_ROOT / "analysis" / "cache"
TZ = ZoneInfo("America/Toronto")

# Where a supplier name might live on an LS product record. The API shape is not
# verified against the live account yet — --probe prints the real keys. Ordered
# most- to least-likely; the first hit wins.
SUPPLIER_PATHS = (("supplier_name",), ("supplier", "name"), ("supplier",))


def today():
    return datetime.now(TZ).date().isoformat()


def dig(record, path):
    """Walk a tuple path through nested dicts, returning None on any miss."""
    cur = record
    for key in path:
        if not isinstance(cur, dict):
            return None
        cur = cur.get(key)
        if cur is None:
            return None
    return cur


def supplier_of(record):
    """The supplier name on an LS product, or None if it carries none."""
    for path in SUPPLIER_PATHS:
        value = dig(record, path)
        if isinstance(value, str) and value.strip():
            return value.strip()
    return None


def probe(client, page_size):
    """Fetch one page and describe it. Writes nothing, caches nothing.

    This exists because the pagination envelope and the product field names are
    configured, not verified — everything downstream assumes them. Run this first
    against a real account and reconcile platform-settings/lightspeed.json with
    what it prints.
    """
    body = client.get(client.products_path(), {client.page_size_param: page_size})
    print("=== response envelope ===")
    if isinstance(body, dict):
        for key, value in body.items():
            if key == "data" and isinstance(value, list):
                print(f"  {key}: list of {len(value)}")
            else:
                print(f"  {key}: {json.dumps(value)[:200]}")
    else:
        print(f"  (bare {type(body).__name__}, length {len(body) if hasattr(body, '__len__') else '?'})")

    records = body.get("data", body) if isinstance(body, dict) else body
    if not isinstance(records, list) or not records:
        print("\nNo records returned — cannot describe a product.")
        return 1

    first = records[0]
    print(f"\n=== first product: {len(first)} fields ===")
    for key in sorted(first):
        print(f"  {key}: {json.dumps(first[key])[:120]}")

    cfg = client.cfg["api"]["pagination"]
    print("\n=== reconcile against platform-settings/lightspeed.json ===")
    print(f"  cursor_field   {cfg['cursor_field']!r}: "
          f"{'present' if cfg['cursor_field'] in first else 'MISSING on the record'}")
    vb = body.get("version") if isinstance(body, dict) else None
    print(f"  version block  : {'present -> ' + json.dumps(vb) if vb else 'absent (cursor comes from records)'}")
    print(f"  supplier name  : {supplier_of(first)!r} via {[p for p in SUPPLIER_PATHS if dig(first, p)]}")
    print(f"\n  requests made  : {client.request_count}")
    return 0


def walk(client, max_pages=None):
    products = list(client.paginate(client.products_path(), max_pages=max_pages))
    return products


def cache_path(date):
    return CACHE_DIR / f"lightspeed-products-{date}.json"


def load_or_walk(client, date, refresh, max_pages):
    path = cache_path(date)
    if path.exists() and not refresh:
        cached = json.loads(path.read_text())
        print(f"cache hit: {path.relative_to(REPO_ROOT)} "
              f"({len(cached['products'])} products, walked {cached['walked_at']}) "
              f"— pass --refresh to re-walk", file=sys.stderr)
        return cached["products"], True

    products = walk(client, max_pages=max_pages)
    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({
        "walked_at": datetime.now(TZ).isoformat(),
        "api_version": client.api_version,
        "stats": client.stats(),
        "products": products,
    }, indent=1) + "\n")
    print(f"walked {len(products)} products in {client.request_count} requests "
          f"-> cached at {path.relative_to(REPO_ROOT)}", file=sys.stderr)
    return products, False


def compare_export(products, export_path):
    """Check a pull against a hand-exported LS workbook.

    The committed exports are the only ground truth available before this script
    has ever run live, so this is the Phase 1 verification: every SKU in the
    export should appear in the pull, carrying the same UUID. Genuine drift since
    the export was taken is expected and is reported, not treated as failure.
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("openpyxl is required for --compare-export (pip install openpyxl)", file=sys.stderr)
        return 2

    # NOT read_only=True. These exports declare a wrong worksheet dimension, and
    # read_only mode trusts the declaration — it reports 1 column and 0 data rows
    # for a 33-column, 267-row sheet. Verified against the two committed exports
    # 2026-09-10. They are under 100 KB, so a normal load costs nothing.
    ws = load_workbook(export_path, data_only=True).active
    rows = ws.iter_rows(values_only=True)
    header = [str(c).strip().lower() if c is not None else "" for c in next(rows)]
    try:
        sku_col, id_col = header.index("sku"), header.index("id")
    except ValueError:
        print(f"export {export_path} has no 'sku'/'id' column; header was {header[:12]}",
              file=sys.stderr)
        return 2

    export = {}
    for row in rows:
        sku = row[sku_col]
        if sku is None or str(sku).strip() == "":
            continue
        export[str(sku).strip()] = (str(row[id_col]).strip() if row[id_col] else None)

    pulled = {}
    for p in products:
        sku = p.get("sku")
        if sku is not None and str(sku).strip():
            pulled[str(sku).strip()] = p.get("id")

    missing = sorted(s for s in export if s not in pulled)
    mismatched = sorted(s for s, uuid in export.items()
                        if s in pulled and uuid and pulled[s] and uuid != pulled[s])

    print(f"export {Path(export_path).name}: {len(export)} SKUs")
    print(f"pull:   {len(pulled)} SKUs")
    print(f"  matched with identical UUID : {len(export) - len(missing) - len(mismatched)}")
    print(f"  in export, absent from pull : {len(missing)}")
    print(f"  UUID disagreement           : {len(mismatched)}")
    for s in missing[:20]:
        print(f"    MISSING    {s}")
    for s in mismatched[:20]:
        print(f"    MISMATCH   {s}: export={export[s]} pull={pulled[s]}")
    if missing or mismatched:
        print("\nMissing SKUs may be genuine drift (products deleted or renamed since "
              "the export). A UUID disagreement is not drift — it means the pull is "
              "reading the wrong field. Investigate those first.")
    return 1 if mismatched else 0


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0],
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--supplier", action="append", default=[],
                    help="Filter to this LS supplier_name. Repeatable. Matched "
                         "case-insensitively; the catalogue is still walked in full.")
    ap.add_argument("--out", type=Path,
                    help="Output path. Default ingest/<today>/lightspeed-products.json")
    ap.add_argument("--refresh", action="store_true",
                    help="Re-walk the catalogue even if today's cache exists.")
    ap.add_argument("--probe", action="store_true",
                    help="Fetch one page, describe the envelope and a product, exit. "
                         "Writes and caches nothing. Run this first.")
    ap.add_argument("--compare-export", type=Path,
                    help="Reconcile the pull against a hand-exported LS .xlsx.")
    ap.add_argument("--max-pages", type=int, help="Stop after N pages (testing).")
    ap.add_argument("--page-size", type=int, help="Override api.pagination.page_size.")
    ap.add_argument("--verbose", action="store_true", help="Log each page to stderr.")
    args = ap.parse_args()

    cfg = load_config()
    try:
        client = LightspeedClient(config=cfg, verbose=args.verbose)
    except LightspeedError as e:
        print(f"error: {e}", file=sys.stderr)
        return 2

    page_size = args.page_size or cfg["api"]["pagination"]["page_size"]

    try:
        if args.probe:
            return probe(client, page_size)

        date = today()
        products, from_cache = load_or_walk(client, date, args.refresh, args.max_pages)
    except LightspeedError as e:
        print(f"error: {e}", file=sys.stderr)
        return 2

    if args.compare_export:
        return compare_export(products, args.compare_export)

    wanted = {s.strip().lower() for s in args.supplier}
    if wanted:
        selected = [p for p in products
                    if (supplier_of(p) or "").lower() in wanted]
    else:
        selected = products

    suppliers = {}
    no_supplier = 0
    for p in selected:
        name = supplier_of(p)
        if name is None:
            no_supplier += 1
        else:
            suppliers[name] = suppliers.get(name, 0) + 1

    out = args.out or (REPO_ROOT / "ingest" / today() / "lightspeed-products.json")
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps({
        "source": "lightspeed",
        "pulled_at": datetime.now(TZ).isoformat(),
        "api_version": client.api_version,
        "from_cache": from_cache,
        "supplier_filter": sorted(wanted) or None,
        "counts": {
            "catalogue_total": len(products),
            "selected": len(selected),
            "without_supplier": no_supplier,
            "by_supplier": dict(sorted(suppliers.items())),
        },
        "stats": client.stats(),
        "products": selected,
    }, indent=1) + "\n")

    print(f"{out.relative_to(REPO_ROOT)} — {len(selected)} of {len(products)} products"
          + (f", filtered to {sorted(wanted)}" if wanted else "")
          + (f", {no_supplier} carry no supplier" if no_supplier else ""))
    if wanted and not selected:
        print(f"  no products matched. Suppliers present in the catalogue: "
              f"{sorted({supplier_of(p) for p in products if supplier_of(p)})}", file=sys.stderr)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
