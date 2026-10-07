#!/usr/bin/env python3
"""Bring a HikePOS product export up to date from the Airtable catalogue (read-only).

    python3 scripts/hike_update.py --supplier FAW \
        --hike-export Products.xlsx \
        --airtable-snapshot ingest/<date>/faw_airtable_snapshot.json \
        --ls-snapshot ingest/<date>/lightspeed-products.json \
        --out ingest/<date>/hike_faw_products_<date>.xlsx \
        --report ingest/<date>/hike_faw_report_<date>.csv

Writes two files and touches no platform. The .xlsx is the Hike export itself, edited in place
(same sheet, same columns, same cell types) with new products appended, ready for Hike's
Products -> IMPORT. The .csv says what changed on which row and why.

Hike has no product id: SKU is the identity, and editing a product's SKU or Product Name in an
import makes a new product. So an existing row keeps Name, SKU and Barcode exactly as exported,
and only price, cost, active state and the description move. Registry: platform-settings/hike.json.
Method: methods/hike-pos.md.
"""

from __future__ import annotations

import argparse
import collections
import csv
import difflib
import json
import re
import sys
from datetime import date
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
REGISTRY = REPO / "platform-settings" / "hike.json"

# Master Flooring Catalogue field ids -> short keys (ids, not names: names get renamed)
AT_FIELDS = {
    "fldx3byCOht5HbKmH": "sku", "fldPSbURaP3wp0Tce": "name", "fldL8miCdB55TWmfJ": "coll",
    "fldvWLfeSvsMnd1vL": "species", "fldb9VH0GrYL2y8OK": "width", "fldnx613WYktJTVpT": "length",
    "fldEevOqBS3QngAGu": "thick", "fldAxs1u4wam2TuNe": "profile", "flduJqOhE4YuX3Bmq": "cost",
    "fldpH6d7GlytLLh7Q": "retail", "fldB21dDwkG60NEAS": "promo", "fldluA0eeTCwfton7": "promoend",
    "fld67650y8QClqoMc": "eff", "fldFXyeHxIVYjWM4R": "box", "fld0sft4ZRTMHt5Hi": "stock",
    "fldPsKISM495IjaWU": "active", "fld8Ra8sQb3SgVD7R": "rep", "fldr4NgCQg3vOeNt1": "repend",
}
SNAPSHOT_FIELD_IDS = list(AT_FIELDS)

# Hike SKU species code (FAW.E.<code>.…) -> a word the Airtable name or collection must contain
SPECIES_CODE = {"H": "hickory", "O": "oak", "0": "oak", "M": "maple", "W": "walnut", "EW": "walnut"}
PROMO_RE = re.compile(r"(Promotion: \$ / )([\d.]+)(=70c)")
LEADING_DATE_RE = re.compile(r"^(\s*(?:<p>)?\s*)(\d{1,4}[/-]\d{1,2}[/-]\d{2,4})(?=\s+-)")
PRICE_MARKER_RE = re.compile(r"^\((?:P|R)\b[^)]*\)\s*")


# ── inputs ────────────────────────────────────────────────────────────────────

def flatten_airtable(raw: dict) -> list[dict]:
    """MCP list_records_for_table output -> [{sku, name, …}] (select values reduced to names)."""
    out = []
    for rec in raw["records"]:
        o = {"rec": rec["id"]}
        for fid, v in (rec.get("cellValuesByFieldId") or rec.get("fields") or {}).items():
            if isinstance(v, dict) and "name" in v:
                v = v["name"]
            o[AT_FIELDS.get(fid, fid)] = v
        if o.get("sku"):
            out.append(o)
    return sorted(out, key=lambda o: o["sku"])


def read_hike(path: Path):
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    rows = []
    for rownum, cells in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        r = dict(zip(header, cells))
        if any(v not in (None, "") for v in r.values()):
            r["_row"] = rownum
            rows.append(r)
    return wb, ws, header, rows


# ── parsing and matching ──────────────────────────────────────────────────────

def norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]", "", s.lower().replace("&", "and"))


def parse_hike_name(name: str | None) -> dict:
    """Category, colour, box size and width from a Hike name such as
    '(C) FAWENG - NAF - Hickory (Bronze) T&G | #FAW.E.H.Bro.6.1 |  - 6.5" x 18mm x … - 19.18sf/b'."""
    name = name or ""
    out = {"cat": None, "colour": None, "box": None, "width": None}
    m = re.search(r"FAW(ENG|HAR|LAM|VIN)", name)
    if m:
        out["cat"] = {"ENG": "ENG", "HAR": "HWD", "LAM": "LAM", "VIN": "VIN"}[m.group(1)]
    m = re.search(r"\(([^()]*)\)\s*(?:Click|T&G)?\s*\|", name)
    if m:
        out["colour"] = m.group(1).strip() or None
    m = re.search(r"([\d.]+)\s*sf/b", name, re.I) or re.search(r"([\d.]+)\s*SQFT/BOX", name, re.I)
    if m:
        try:
            out["box"] = float(m.group(1).rstrip("."))
        except ValueError:
            pass
    m = re.search(r'\|\s*-?\s*(\d+(?:\.\d+)?)(?:\s+(\d)/(\d))?"', name)
    if m:
        w = float(m.group(1))
        if m.group(2):
            w += int(m.group(2)) / int(m.group(3))
        out["width"] = w
    else:
        m = re.search(r'\((\d+(?:\.\d+)?)" wide', name)
        out["width"] = float(m.group(1)) if m else None
    return out


def at_cat(o: dict) -> str:
    c = o["sku"].split("-")[0]
    return "VIN" if c in ("LVP", "LVT") else c


def at_colour(o: dict) -> str | None:
    return o["name"].split("—")[-1].strip() if "—" in (o.get("name") or "") else None


def _species_ok(hike_sku: str, o: dict) -> bool:
    m = re.match(r"FAW\.[EH]\.(\w+?)\.", hike_sku or "")
    word = SPECIES_CODE.get(m.group(1)) if m else None
    return word is None or word in (o["name"] + " " + (o.get("coll") or "")).lower()


def _kind_ok(hike_name: str, o: dict) -> bool:
    tile = any(t in (hike_name or "") for t in ("Tile", "609", "914"))
    if o["sku"].startswith("LVT"):
        return tile
    if o["sku"].startswith("LVP"):
        return not tile
    return True


def _width_ok(p: dict, o: dict) -> bool:
    if p["width"] is None or o.get("width") is None or o["sku"].startswith("LVT"):
        return True
    return abs(p["width"] - o["width"]) <= 0.5


def match(hike_rows: list[dict], at_records: list[dict], cfg: dict):
    """Return (matches {hike index: airtable sku}, primary {airtable sku: hike index}, notes).

    A Hike row matches when it names the same colour in the same category and nothing in its
    species, plank-or-tile kind or width contradicts the record. Several Hike rows may be the
    same Airtable product (Hike carries duplicates); all of them get the new cost and price,
    and the primary one (active, same box size, a structured FAW. SKU) also gets Active."""
    at = {o["sku"]: o for o in at_records}
    aliases = cfg.get("colour_aliases", {})
    manual = cfg.get("manual_matches", {})
    matches, notes = {}, collections.defaultdict(list)
    for i, h in enumerate(hike_rows):
        sku = str(h.get("SKU") or "")
        if sku in manual:
            matches[i] = manual[sku]
            notes[i].append("hand match (registry)")
            continue
        p = parse_hike_name(h.get("Name"))
        if p["colour"] in aliases:
            matches[i] = aliases[p["colour"]]
            notes[i].append(f"spelling alias {p['colour']!r}")
            continue
        if not p["cat"] or not p["colour"]:
            continue
        cands = []
        for o in at_records:
            if at_cat(o) != p["cat"] or norm(at_colour(o) or "") != norm(p["colour"]):
                continue
            if not _species_ok(sku, o) or not _kind_ok(h.get("Name"), o):
                continue
            if not _width_ok(p, o):
                notes[i].append(f"width {p['width']} rules out {o['sku']} ({o.get('width')})")
                continue
            cands.append(o["sku"])
        if len(cands) > 1 and p["box"]:
            best = min(abs((at[s].get("box") or 0) - p["box"]) for s in cands)
            cands = [s for s in cands if abs((at[s].get("box") or 0) - p["box"]) - best < 0.01]
        if len(cands) == 1:
            matches[i] = cands[0]
        elif len(cands) > 1:
            notes[i].append("ambiguous: " + ", ".join(cands))
    by_at = collections.defaultdict(list)
    for i, s in matches.items():
        by_at[s].append(i)
    primary = {}
    for s, ids in by_at.items():
        box = at[s].get("box") or 0

        def key(i):
            p = parse_hike_name(hike_rows[i].get("Name"))
            boxd = abs(p["box"] - box) if p["box"] else 9
            return (hike_rows[i].get("Active") != "TRUE", boxd > 0.05,
                    not str(hike_rows[i].get("SKU")).startswith("FAW."), boxd)
        primary[s] = sorted(ids, key=key)[0]
    return matches, primary, notes


# ── values ────────────────────────────────────────────────────────────────────

def effective_cost(o: dict, today: str) -> float | None:
    """The lowest of Cost/unit, an active promo and an active rep rate — what the POS pays now."""
    costs = [o.get("cost")]
    if o.get("promo") is not None and (o.get("promoend") or "") >= today:
        costs.append(o["promo"])
    if o.get("rep") is not None and (not o.get("repend") or o["repend"] >= today):
        costs.append(o["rep"])
    costs = [c for c in costs if c is not None]
    return round(min(costs), 2) if costs else None


def new_retail(o: dict, cost: float, old_cost, old_retail, markup: float) -> float:
    """Flooring: cost + the Hike markup. An accessory keeps the markup it already had."""
    if o["sku"].startswith("ACC") and old_cost is not None and old_retail is not None:
        return round(cost + (old_retail - old_cost), 2)
    return round(cost + markup, 2)


def rewrite_description(desc, eff: str | None, old_cost, new_cost: float, offset: float):
    """Move the leading price date to the Airtable Effective Date, and the 'Promotion: $ / N=70c'
    floor to new cost + offset — only where N was old cost + offset (flooring rows)."""
    if not desc:
        return desc
    out = desc
    if eff:
        out = LEADING_DATE_RE.sub(lambda m: m.group(1) + eff, out, count=1)
    m = PROMO_RE.search(out)
    if m and old_cost is not None and abs(float(m.group(2)) - (old_cost + offset)) < 0.005:
        out = out[:m.start(2)] + f"{new_cost + offset:.2f}" + out[m.end(2):]
    return out


def hike_name_for_new(o: dict, cfg: dict, ls_names: dict) -> str:
    """Hike name for a product new to Hike, in the existing '(C) FAWxxx - NAF - … | #SKU | …'
    form. Taken from the product's Lightspeed name (built from Airtable) with any (P …)/(R …)
    price marker removed; built from Airtable fields when there is no Lightspeed name."""
    prefix, sku = cfg["name_prefix"], o["sku"]
    ls = ls_names.get(sku)
    if ls:
        ls = PRICE_MARKER_RE.sub("", ls)
        head, _, specs = ls.partition(" | ")
        tok, _, desc = head.partition(" - ")
        tok = cfg["lightspeed_name_prefix"].get(tok.strip(), cfg["category_token"][sku.split("-")[0]])
        tail = f" |  - {specs}" if specs else " |"
        return f"{prefix}{tok} - NAF - {desc.strip()} | #{sku}{tail}"
    tok = cfg["category_token"][sku.split("-")[0]]
    dims = " x ".join(x for x in (
        f'{o["width"]:g}"' if o.get("width") else "",
        f'{o["thick"]:g}mm' if o.get("thick") else "",
        o.get("length") or "") if x)
    box = f' - {o["box"]:g}sf/b' if o.get("box") else ""
    profile = f" {o['profile']}" if o.get("profile") in ("Click", "T&G") else ""
    return f"{prefix}{tok} - NAF - {o.get('coll') or ''} ({at_colour(o) or o['name']}){profile} | #{sku} |  - {dims}{box}"


# ── plan and write ────────────────────────────────────────────────────────────

def build(hike_rows, header, at_records, cfg, outlet, ls_names, today):
    at = {o["sku"]: o for o in at_records}
    markup, offset = cfg["markup"], cfg["promo_floor_offset"]
    col = lambda name: f"{outlet}_{name}"
    matches, primary, notes = match(hike_rows, at_records, cfg)
    is_primary = {i for i in primary.values()}
    updates, report = {}, []

    def rep(action, h, o, field="", old="", new="", note=""):
        report.append({"action": action, "hike_sku": h.get("SKU") if h else "", "hike_name": h.get("Name") if h else "",
                       "airtable_sku": o["sku"] if o else "", "airtable_name": o["name"] if o else "",
                       "field": field, "old": old, "new": new, "note": note})

    for i, s in sorted(matches.items()):
        h, o = hike_rows[i], at[s]
        cost = effective_cost(o, today)
        if cost is None:
            rep("flag", h, o, note="Airtable has no cost; row left as exported")
            continue
        old_cost, old_retail = h.get(col("Cost price")), h.get(col("Retail price"))
        retail = new_retail(o, cost, old_cost, old_retail, markup)
        ch = {col("Cost price"): cost, col("Retail price"): retail, col("Price Excluding Tax"): retail}
        desc = rewrite_description(h.get("Description"), o.get("eff"), old_cost, cost, offset)
        if desc != h.get("Description"):
            ch["Description"] = desc
        if i in is_primary:
            ch["Active"] = "TRUE" if o.get("active") and o.get("stock") != "Discontinued" else "FALSE"
        ch = {k: v for k, v in ch.items() if h.get(k) != v}
        updates[i] = ch
        note = "; ".join(notes.get(i, []))
        if i not in is_primary:
            note = (note + "; " if note else "") + f"duplicate of {hike_rows[primary[s]]['SKU']} in Hike: price updated, Active left as is"
        p = parse_hike_name(h.get("Name"))
        if p["box"] and o.get("box") and abs(p["box"] - o["box"]) > 0.05:
            note = (note + "; " if note else "") + f"name says {p['box']}sf/b, Airtable {o['box']} (name not changed: Hike would make a new product)"
        if not ch:
            rep("unchanged", h, o, note=note)
        for k, v in ch.items():
            rep("update", h, o, k, h.get(k), v, note)

    covered = set(matches.values())
    for i, h in enumerate(hike_rows):
        if i not in matches:
            rep("hike_only", h, None, note="; ".join(notes.get(i, [])) or "no Airtable product; left exactly as exported")

    # new products: active Airtable records with no Hike row
    template = next(h for h in hike_rows if str(h.get("SKU", "")).startswith("FAW.") and h.get("Active") == "TRUE")
    sibling_desc = {}
    for s, i in primary.items():
        key = (at_cat(at[s]), at[s].get("coll"))
        if PROMO_RE.search(hike_rows[i].get("Description") or "") and key not in sibling_desc:
            sibling_desc[key] = (hike_rows[i]["Description"], hike_rows[i].get(col("Cost price")))
    new_rows = []
    for o in at_records:
        s = o["sku"]
        if s in covered or s in cfg.get("exclude_new", []) or not o.get("active") or o.get("stock") == "Discontinued":
            continue
        cost = effective_cost(o, today)
        if cost is None:
            rep("flag", None, o, note="new to Hike but Airtable has no cost; not added")
            continue
        cat = s.split("-")[0]
        row = {k: template.get(k) for k in header}
        sib = sibling_desc.get((at_cat(o), o.get("coll")))
        # a sibling in the same collection lends its install/spec notes (collection-level, e.g.
        # every Aquaplus Gold row carries the same two-factories note); stray HTML is dropped
        desc = (re.sub(r"</?p>", "", rewrite_description(sib[0], o.get("eff"), sib[1], cost, offset)).strip() if sib
                else f"{o.get('eff') or ''} -  -  / Promotion: $ / {cost + offset:.2f}=70c")
        row.update({
            "Name": hike_name_for_new(o, cfg, ls_names), "Description": desc, "SKU": s, "Barcode": s,
            "Product type": cfg["product_type"][cat], "Product tag": cfg["product_tag"],
            "Supplier name": cfg["hike_supplier_name"], "Slug": None, "Image URL": None, "Active": "TRUE",
            "Track inventory": "FALSE", "Allow out of stock": "TRUE",
            col("Cost price"): cost, col("Retail price"): round(cost + markup, 2),
            col("Price Excluding Tax"): round(cost + markup, 2), col("Average Cost Price"): cost,
            col("Stock"): 0, col("Stock on hand"): 0, col("Reorder level"): None, col("Reorder value"): None,
        })
        new_rows.append(row)
        rep("new", {"SKU": s, "Name": row["Name"]}, o, note="Airtable SKU used as Hike SKU and Barcode")
    return updates, new_rows, report


def write_outputs(wb, ws, header, hike_rows, updates, new_rows, out: Path, report_rows, report: Path, cfg_never):
    colidx = {h: n + 1 for n, h in enumerate(header)}
    for i, ch in updates.items():
        for k, v in ch.items():
            if k in cfg_never:
                raise SystemExit(f"refusing to change {k} on an existing row ({hike_rows[i]['SKU']})")
            ws.cell(row=hike_rows[i]["_row"], column=colidx[k], value=v)
    nxt = max(h["_row"] for h in hike_rows) + 1
    for r in new_rows:
        for k, v in r.items():
            ws.cell(row=nxt, column=colidx[k], value=v)
        nxt += 1
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    with open(report, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["action", "hike_sku", "hike_name", "airtable_sku", "airtable_name",
                                          "field", "old", "new", "note"])
        w.writeheader()
        for r in report_rows:
            w.writerow({k: "" if v is None else v for k, v in r.items()})


def load_ls_names(path: Path | None) -> dict:
    if not path:
        return {}
    d = json.loads(path.read_text())
    return {p["sku"]: p["name"] for p in d.get("products", d) if p.get("sku") and p.get("name")}


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--supplier", required=True)
    ap.add_argument("--hike-export", type=Path, required=True)
    ap.add_argument("--airtable-snapshot", type=Path, required=True)
    ap.add_argument("--ls-snapshot", type=Path)
    ap.add_argument("--out", type=Path, required=True)
    ap.add_argument("--report", type=Path, required=True)
    ap.add_argument("--today", default=date.today().isoformat())
    a = ap.parse_args(argv)
    try:
        import openpyxl  # noqa: F401
    except ImportError:
        print("openpyxl is required (pip install openpyxl)", file=sys.stderr)
        return 2
    reg = json.loads(REGISTRY.read_text())
    cfg = reg["suppliers"][a.supplier]
    wb, ws, header, hike_rows = read_hike(a.hike_export)
    at_records = flatten_airtable(json.loads(a.airtable_snapshot.read_text()))
    updates, new_rows, report = build(hike_rows, header, at_records, cfg, reg["outlet"],
                                      load_ls_names(a.ls_snapshot), a.today)
    write_outputs(wb, ws, header, hike_rows, updates, new_rows, a.out, report, a.report,
                  set(reg["never_change_on_existing"]) | set(reg["never_touch_columns"]))
    c = collections.Counter(r["action"] for r in report)
    changed = sum(1 for ch in updates.values() if ch)
    print(f"Hike rows: {len(hike_rows)}  matched: {len(updates)} ({changed} changed)  "
          f"new: {len(new_rows)}  hike-only: {c['hike_only']}  flags: {c['flag']}")
    print(f"wrote {a.out}\nwrote {a.report}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
